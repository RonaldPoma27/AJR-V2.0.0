import csv
import io
import re
import unicodedata
from urllib.parse import parse_qs, urljoin, urlsplit
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated, Any
from uuid import uuid4

from fastapi import APIRouter, File, Query, UploadFile
from starlette.datastructures import UploadFile as StarletteUploadFile
from sqlalchemy import and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import DBSession, StaffUser
from app.core.exceptions import AppError
from app.models import User, UserRole
from app.models.agile import (
    AgileMembership, AgileProject, AgileSprint, AgileSprintItem,
    AgileWorkItem, AgileWorkItemComment, AgileWorkItemHistory,
)
from app.schemas.agile import (
    GoogleSheetImportRequest,
    ImportPreview, ImportPreviewRow, ImportResult, MemberCreate, MemberRead,
    ProjectCreate, ProjectRead, ProjectUpdate, SprintComplete, SprintCreate,
    SprintRead, StaffOption, WorkItemCommentCreate, WorkItemCommentRead, WorkItemCreate, WorkItemHistoryRead, WorkItemRead, WorkItemReorder, WorkItemUpdate,
)

router = APIRouter(prefix="/agile", tags=["agile-projects"])
PROJECT_ROLES = {"SCRUM_MASTER", "PRODUCT_OWNER", "DEVELOPER", "OBSERVER"}
PLAN_ROLES = {"ADMIN", "SCRUM_MASTER", "PRODUCT_OWNER"}
EDIT_ROLES = PLAN_ROLES | {"DEVELOPER"}
STATUSES = {"TODO", "IN_PROGRESS", "IN_REVIEW", "DONE", "BLOCKED"}
ITEM_TYPES = {"EPIC", "STORY", "TASK", "BUG", "SUBTASK"}
PRIORITIES = {"LOW", "MEDIUM", "HIGH", "URGENT"}
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000


def _role_name(user: User) -> str:
    return user.role.value if hasattr(user.role, "value") else str(user.role)


def _can(role: str, allowed: set[str]) -> bool:
    return role in allowed


async def _project_access(db: AsyncSession, project_id: int, user: User) -> tuple[AgileProject, str]:
    project = await db.get(AgileProject, project_id)
    if project is None:
        raise AppError(404, "Proyecto ágil no encontrado.")
    if user.role == UserRole.ADMIN:
        return project, "ADMIN"
    membership = await db.scalar(
        select(AgileMembership).where(
            AgileMembership.project_id == project_id,
            AgileMembership.user_id == user.id,
        )
    )
    if membership is None:
        # El proyecto no se revela a técnicos que no pertenezcan a él.
        raise AppError(404, "Proyecto ágil no encontrado.")
    return project, membership.role


async def _require_role(db: AsyncSession, project_id: int, user: User, allowed: set[str]) -> tuple[AgileProject, str]:
    project, role = await _project_access(db, project_id, user)
    if not _can(role, allowed):
        raise AppError(403, "Tu rol en este proyecto no permite realizar esta acción.")
    if project.status == "ARCHIVED" and role != "ADMIN":
        raise AppError(409, "El proyecto está archivado y es de solo lectura.")
    return project, role


async def _membership(db: AsyncSession, project_id: int, user_id: int) -> AgileMembership | None:
    return await db.scalar(select(AgileMembership).where(
        AgileMembership.project_id == project_id, AgileMembership.user_id == user_id
    ))


async def _validate_assignee(db: AsyncSession, project_id: int, user_id: int | None) -> None:
    if user_id is None:
        return
    user = await db.get(User, user_id)
    if user is None or user.role == UserRole.USER:
        raise AppError(422, "El responsable debe ser una cuenta ADMIN o TECHNICIAN.")
    if await _membership(db, project_id, user_id) is None:
        raise AppError(422, "El responsable debe formar parte del equipo del proyecto.")


async def _validate_sprint(db: AsyncSession, project_id: int, sprint_id: int | None) -> AgileSprint | None:
    if sprint_id is None:
        return None
    sprint = await db.get(AgileSprint, sprint_id)
    if sprint is None or sprint.project_id != project_id:
        raise AppError(422, "El sprint indicado no pertenece a este proyecto.")
    if sprint.status == "COMPLETED":
        raise AppError(409, "No se pueden agregar tarjetas a un sprint completado.")
    return sprint


async def _validate_epic(db: AsyncSession, project_id: int, item_type: str, epic_id: int | None, item_id: int | None = None) -> AgileWorkItem | None:
    if item_type == "EPIC":
        if epic_id is not None:
            raise AppError(422, "Una épica no puede depender de otra épica.")
        return None
    if epic_id is None:
        return None
    if item_id is not None and epic_id == item_id:
        raise AppError(422, "Una tarjeta no puede ser su propia épica.")
    epic = await db.get(AgileWorkItem, epic_id)
    if epic is None or epic.project_id != project_id or epic.item_type != "EPIC" or epic.is_archived:
        raise AppError(422, "La épica seleccionada no existe o no pertenece a este proyecto.")
    return epic


async def _add_history(db: AsyncSession, item: AgileWorkItem, actor: User, field: str, old: Any, new: Any) -> None:
    if old == new:
        return
    db.add(AgileWorkItemHistory(
        project_id=item.project_id,
        work_item_id=item.id,
        actor_id=actor.id,
        field_name=field,
        old_value=None if old is None else str(old),
        new_value=None if new is None else str(new),
    ))


async def _record_sprint_item(db: AsyncSession, sprint_id: int | None, item_id: int) -> None:
    if sprint_id is None:
        return
    exists = await db.scalar(select(AgileSprintItem.id).where(
        AgileSprintItem.sprint_id == sprint_id,
        AgileSprintItem.work_item_id == item_id,
    ))
    if exists is None:
        db.add(AgileSprintItem(sprint_id=sprint_id, work_item_id=item_id))


async def _item_read(db: AsyncSession, item: AgileWorkItem) -> dict[str, Any]:
    assignee = await db.get(User, item.assignee_id) if item.assignee_id else None
    epic = await db.get(AgileWorkItem, item.project_epic_id) if item.project_epic_id else None
    sprint = await db.get(AgileSprint, item.sprint_id) if item.sprint_id else None
    return {
        "id": item.id,
        "project_id": item.project_id,
        "key": item.key,
        "title": item.title,
        "description": item.description,
        "acceptance_criteria": item.acceptance_criteria,
        "item_type": item.item_type,
        "status": item.status,
        "priority": item.priority,
        "assigned_role": item.assigned_role,
        "estimate_points": item.estimate_points,
        "order_index": item.order_index,
        "epic_id": item.project_epic_id,
        "epic_key": epic.key if epic else None,
        "epic_title": epic.title if epic else None,
        "assignee_id": item.assignee_id,
        "assignee_name": assignee.full_name if assignee else None,
        "assignee_email": assignee.email if assignee else None,
        "sprint_id": item.sprint_id,
        "sprint_name": sprint.name if sprint else None,
        "due_date": item.due_date,
        "created_by_id": item.created_by_id,
        "is_archived": item.is_archived,
        "created_at": item.created_at,
        "updated_at": item.updated_at,
        "completed_at": item.completed_at,
    }


async def _item_reads(db: AsyncSession, items: list[AgileWorkItem]) -> list[dict[str, Any]]:
    """Serializa una colección completa con consultas por lote (evita N+1 en tableros grandes)."""
    if not items:
        return []
    user_ids = {item.assignee_id for item in items if item.assignee_id is not None}
    epic_ids = {item.project_epic_id for item in items if item.project_epic_id is not None}
    sprint_ids = {item.sprint_id for item in items if item.sprint_id is not None}
    users = {u.id: u for u in (await db.scalars(select(User).where(User.id.in_(user_ids)))).all()} if user_ids else {}
    epics = {e.id: e for e in (await db.scalars(select(AgileWorkItem).where(AgileWorkItem.id.in_(epic_ids)))).all()} if epic_ids else {}
    sprints = {s.id: s for s in (await db.scalars(select(AgileSprint).where(AgileSprint.id.in_(sprint_ids)))).all()} if sprint_ids else {}
    result: list[dict[str, Any]] = []
    for item in items:
        assignee = users.get(item.assignee_id)
        epic = epics.get(item.project_epic_id)
        sprint = sprints.get(item.sprint_id)
        result.append({
            "id": item.id, "project_id": item.project_id, "key": item.key, "title": item.title,
            "description": item.description, "acceptance_criteria": item.acceptance_criteria,
            "item_type": item.item_type, "status": item.status, "priority": item.priority,
            "assigned_role": item.assigned_role, "estimate_points": item.estimate_points, "order_index": item.order_index,
            "epic_id": item.project_epic_id, "epic_key": epic.key if epic else None,
            "epic_title": epic.title if epic else None, "assignee_id": item.assignee_id,
            "assignee_name": assignee.full_name if assignee else None,
            "assignee_email": assignee.email if assignee else None,
            "sprint_id": item.sprint_id, "sprint_name": sprint.name if sprint else None,
            "due_date": item.due_date, "created_by_id": item.created_by_id,
            "is_archived": item.is_archived, "created_at": item.created_at,
            "updated_at": item.updated_at, "completed_at": item.completed_at,
        })
    return result


async def _sprint_read(db: AsyncSession, sprint: AgileSprint) -> dict[str, Any]:
    # Se calcula contra la tabla de inclusión histórica, no contra sprint_id actual: al cerrar
    # un sprint las tarjetas salen de él o se transfieren y no deben desaparecer de sus métricas.
    counts = (await db.execute(
        select(
            func.count(AgileSprintItem.work_item_id),
            func.coalesce(func.sum(AgileWorkItem.estimate_points), 0),
        ).join(AgileWorkItem, AgileWorkItem.id == AgileSprintItem.work_item_id)
        .where(AgileSprintItem.sprint_id == sprint.id)
    )).one()
    done = (await db.execute(
        select(
            func.count(AgileSprintItem.work_item_id),
            func.coalesce(func.sum(AgileWorkItem.estimate_points), 0),
        ).join(AgileWorkItem, AgileWorkItem.id == AgileSprintItem.work_item_id)
        .where(
            AgileSprintItem.sprint_id == sprint.id,
            or_(AgileSprintItem.outcome == "COMPLETED", and_(AgileSprintItem.outcome.is_(None), AgileWorkItem.status == "DONE")),
        )
    )).one()
    return {
        "id": sprint.id, "project_id": sprint.project_id, "name": sprint.name,
        "goal": sprint.goal, "start_date": sprint.start_date, "end_date": sprint.end_date,
        "status": sprint.status, "created_by_id": sprint.created_by_id,
        "completed_at": sprint.completed_at, "created_at": sprint.created_at,
        "updated_at": sprint.updated_at, "total_items": counts[0], "total_points": int(counts[1] or 0),
        "completed_items": done[0], "completed_points": int(done[1] or 0),
    }


async def _project_read(db: AsyncSession, project: AgileProject, role: str) -> dict[str, Any]:
    total = await db.scalar(select(func.count(AgileWorkItem.id)).where(
        AgileWorkItem.project_id == project.id, AgileWorkItem.is_archived.is_(False)
    )) or 0
    done = await db.scalar(select(func.count(AgileWorkItem.id)).where(
        AgileWorkItem.project_id == project.id, AgileWorkItem.is_archived.is_(False), AgileWorkItem.status == "DONE"
    )) or 0
    sprint = await db.scalar(select(AgileSprint).where(
        AgileSprint.project_id == project.id, AgileSprint.status == "ACTIVE"
    ))
    return {
        "id": project.id, "key": project.key, "name": project.name,
        "description": project.description, "status": project.status,
        "sprint_length_weeks": project.sprint_length_weeks, "owner_id": project.owner_id,
        "created_at": project.created_at, "updated_at": project.updated_at,
        "my_role": role, "total_items": total, "completed_items": done,
        "active_sprint_name": sprint.name if sprint else None,
    }


@router.get("/users", response_model=list[StaffOption])
async def list_staff_users(db: DBSession, _: StaffUser):
    users = (await db.scalars(select(User).where(User.role.in_([UserRole.ADMIN, UserRole.TECHNICIAN])).order_by(User.last_name, User.first_name, User.email))).all()
    return [{"id": u.id, "email": u.email, "full_name": u.full_name, "global_role": _role_name(u)} for u in users]


@router.get("/projects", response_model=list[ProjectRead])
async def list_projects(db: DBSession, actor: StaffUser):
    if actor.role == UserRole.ADMIN:
        projects = (await db.scalars(select(AgileProject).order_by(AgileProject.updated_at.desc(), AgileProject.id.desc()))).all()
        roles = {p.id: "ADMIN" for p in projects}
    else:
        memberships = (await db.execute(
            select(AgileProject, AgileMembership.role)
            .join(AgileMembership, AgileMembership.project_id == AgileProject.id)
            .where(AgileMembership.user_id == actor.id)
            .order_by(AgileProject.updated_at.desc(), AgileProject.id.desc())
        )).all()
        projects = [row[0] for row in memberships]
        roles = {row[0].id: row[1] for row in memberships}
    return [await _project_read(db, project, roles[project.id]) for project in projects]


@router.post("/projects", response_model=ProjectRead, status_code=201)
async def create_project(payload: ProjectCreate, db: DBSession, actor: StaffUser):
    duplicate = await db.scalar(select(AgileProject.id).where(AgileProject.key == payload.key))
    if duplicate is not None:
        raise AppError(409, "Ya existe un proyecto con esa clave.")
    project = AgileProject(
        key=payload.key, name=payload.name.strip(), description=payload.description.strip(),
        sprint_length_weeks=payload.sprint_length_weeks, owner_id=actor.id,
    )
    db.add(project)
    await db.flush()
    db.add(AgileMembership(project_id=project.id, user_id=actor.id, role="SCRUM_MASTER"))
    await db.commit()
    return await _project_read(db, project, "ADMIN" if actor.role == UserRole.ADMIN else "SCRUM_MASTER")


@router.get("/projects/{project_id}", response_model=ProjectRead)
async def get_project(project_id: int, db: DBSession, actor: StaffUser):
    project, role = await _project_access(db, project_id, actor)
    return await _project_read(db, project, role)


@router.patch("/projects/{project_id}", response_model=ProjectRead)
async def update_project(project_id: int, payload: ProjectUpdate, db: DBSession, actor: StaffUser):
    project, role = await _require_role(db, project_id, actor, PLAN_ROLES)
    for name, value in payload.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(project, name, value.strip() if name in {"name", "description"} else value)
    await db.commit()
    return await _project_read(db, project, role)


@router.get("/projects/{project_id}/members", response_model=list[MemberRead])
async def list_members(project_id: int, db: DBSession, actor: StaffUser):
    await _project_access(db, project_id, actor)
    rows = (await db.execute(
        select(AgileMembership, User).join(User, User.id == AgileMembership.user_id)
        .where(AgileMembership.project_id == project_id)
        .order_by(AgileMembership.joined_at, User.email)
    )).all()
    return [{"id": m.id, "project_id": m.project_id, "user_id": u.id, "email": u.email,
             "full_name": u.full_name, "role": m.role, "joined_at": m.joined_at} for m, u in rows]


@router.post("/projects/{project_id}/members", response_model=MemberRead, status_code=201)
async def add_member(project_id: int, payload: MemberCreate, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    user = await db.get(User, payload.user_id)
    if user is None or user.role == UserRole.USER:
        raise AppError(422, "Solo se pueden incorporar cuentas ADMIN o TECHNICIAN al equipo ágil.")
    member = await _membership(db, project_id, payload.user_id)
    if member is not None:
        member.role = payload.role
    else:
        member = AgileMembership(project_id=project_id, user_id=payload.user_id, role=payload.role)
        db.add(member)
    await db.commit()
    await db.refresh(member)
    return {"id": member.id, "project_id": project_id, "user_id": user.id,
            "email": user.email, "full_name": user.full_name, "role": member.role,
            "joined_at": member.joined_at}


@router.patch("/projects/{project_id}/members/{member_id}", response_model=MemberRead)
async def update_member(project_id: int, member_id: int, payload: MemberCreate, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    member = await db.get(AgileMembership, member_id)
    if member is None or member.project_id != project_id:
        raise AppError(404, "Integrante no encontrado.")
    user = await db.get(User, payload.user_id)
    if user is None or user.role == UserRole.USER:
        raise AppError(422, "Solo se pueden incorporar cuentas ADMIN o TECHNICIAN al equipo ágil.")
    other = await _membership(db, project_id, payload.user_id)
    if other is not None and other.id != member.id:
        raise AppError(409, "La cuenta ya forma parte del equipo.")
    member.user_id, member.role = payload.user_id, payload.role
    await db.commit()
    return {"id": member.id, "project_id": project_id, "user_id": user.id,
            "email": user.email, "full_name": user.full_name, "role": member.role,
            "joined_at": member.joined_at}


@router.delete("/projects/{project_id}/members/{member_id}", status_code=204)
async def remove_member(project_id: int, member_id: int, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    member = await db.get(AgileMembership, member_id)
    if member is None or member.project_id != project_id:
        raise AppError(404, "Integrante no encontrado.")
    if member.user_id == actor.id and member.role == "SCRUM_MASTER":
        another = await db.scalar(select(AgileMembership.id).where(
            AgileMembership.project_id == project_id, AgileMembership.id != member.id,
            AgileMembership.role == "SCRUM_MASTER"
        ))
        if another is None and actor.role != UserRole.ADMIN:
            raise AppError(409, "Asigná otro Scrum Master antes de quitarte del proyecto.")
    assigned_items = (await db.scalars(select(AgileWorkItem).where(
        AgileWorkItem.project_id == project_id, AgileWorkItem.assignee_id == member.user_id,
        AgileWorkItem.is_archived.is_(False)
    ))).all()
    for item in assigned_items:
        await _add_history(db, item, actor, "assignee_id", member.user_id, None)
        item.assignee_id = None
    await db.delete(member)
    await db.commit()
    return None


@router.post("/projects/{project_id}/backlog/reorder", response_model=list[WorkItemRead])
async def reorder_backlog(project_id: int, payload: WorkItemReorder, db: DBSession, actor: StaffUser):
    """Guarda el orden prioritario del backlog como una sola operación."""
    await _require_role(db, project_id, actor, EDIT_ROLES)
    backlog_items = list((await db.scalars(
        select(AgileWorkItem).where(
            AgileWorkItem.project_id == project_id,
            AgileWorkItem.sprint_id.is_(None),
            AgileWorkItem.item_type != "EPIC",
            AgileWorkItem.is_archived.is_(False),
        ).order_by(AgileWorkItem.order_index, AgileWorkItem.id)
    )).all())
    current_ids = {item.id for item in backlog_items}
    if len(payload.item_ids) != len(backlog_items) or set(payload.item_ids) != current_ids:
        raise AppError(409, "El backlog cambió mientras lo ordenabas. Actualizá la vista e intentá nuevamente.")
    by_id = {item.id: item for item in backlog_items}
    for order_index, item_id in enumerate(payload.item_ids):
        item = by_id[item_id]
        previous = item.order_index
        item.order_index = order_index
        await _add_history(db, item, actor, "order_index", previous, order_index)
    await db.commit()
    return await _item_reads(db, [by_id[item_id] for item_id in payload.item_ids])


@router.get("/projects/{project_id}/items", response_model=list[WorkItemRead])
async def list_items(
    project_id: int, db: DBSession, actor: StaffUser,
    item_type: Annotated[str | None, Query(alias="type")] = None,
    sprint_id: int | None = None,
    backlog: bool = False,
    include_archived: bool = False,
):
    await _project_access(db, project_id, actor)
    stmt = select(AgileWorkItem).where(AgileWorkItem.project_id == project_id)
    if not include_archived:
        stmt = stmt.where(AgileWorkItem.is_archived.is_(False))
    if item_type:
        if item_type not in ITEM_TYPES:
            raise AppError(422, "Tipo de tarjeta inválido.")
        stmt = stmt.where(AgileWorkItem.item_type == item_type)
    if sprint_id is not None:
        stmt = stmt.where(AgileWorkItem.sprint_id == sprint_id)
    elif backlog:
        stmt = stmt.where(AgileWorkItem.sprint_id.is_(None))
    items = list((await db.scalars(stmt.order_by(AgileWorkItem.order_index, AgileWorkItem.id).limit(5000))).all())
    return await _item_reads(db, items)


@router.post("/projects/{project_id}/items", response_model=WorkItemRead, status_code=201)
async def create_item(project_id: int, payload: WorkItemCreate, db: DBSession, actor: StaffUser):
    project, _ = await _require_role(db, project_id, actor, EDIT_ROLES)
    title = payload.title.strip()
    if not title:
        raise AppError(422, "El título no puede quedar vacío.")
    epic = await _validate_epic(db, project_id, payload.item_type, payload.epic_id)
    await _validate_assignee(db, project_id, payload.assignee_id)
    await _validate_sprint(db, project_id, payload.sprint_id)
    item = AgileWorkItem(
        project_id=project_id, key=f"TMP-{uuid4().hex[:12].upper()}", title=title, description=payload.description.strip(),
        acceptance_criteria=payload.acceptance_criteria.strip(), item_type=payload.item_type,
        status=payload.status, priority=payload.priority, assigned_role=payload.assigned_role,
        estimate_points=payload.estimate_points, order_index=payload.order_index,
        project_epic_id=epic.id if epic else None,
        assignee_id=payload.assignee_id, sprint_id=payload.sprint_id, due_date=payload.due_date,
        created_by_id=actor.id, completed_at=datetime.now(timezone.utc) if payload.status == "DONE" else None,
    )
    db.add(item)
    await db.flush()
    item.key = f"{project.key}-{item.id:04d}"
    db.add(AgileWorkItemHistory(project_id=project_id, work_item_id=item.id, actor_id=actor.id,
                                field_name="created", old_value=None, new_value=item.key))
    await _record_sprint_item(db, item.sprint_id, item.id)
    await db.commit()
    return await _item_read(db, item)


@router.patch("/projects/{project_id}/items/{item_id}", response_model=WorkItemRead)
async def update_item(project_id: int, item_id: int, payload: WorkItemUpdate, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, EDIT_ROLES)
    item = await db.get(AgileWorkItem, item_id)
    if item is None or item.project_id != project_id or item.is_archived:
        raise AppError(404, "Tarjeta no encontrada.")
    data = payload.model_dump(exclude_unset=True)
    if "title" in data:
        data["title"] = (data["title"] or "").strip()
        if not data["title"]:
            raise AppError(422, "El título no puede quedar vacío.")
    next_type = data.get("item_type", item.item_type)
    next_epic_id = data.get("epic_id", item.project_epic_id)
    epic = await _validate_epic(db, project_id, next_type, next_epic_id, item.id)
    if "epic_id" in data or "item_type" in data:
        data["project_epic_id"] = epic.id if epic else None
        data.pop("epic_id", None)
    if "assignee_id" in data:
        await _validate_assignee(db, project_id, data["assignee_id"])
    if "sprint_id" in data:
        await _validate_sprint(db, project_id, data["sprint_id"])
    before = {
        "title": item.title, "description": item.description, "acceptance_criteria": item.acceptance_criteria,
        "item_type": item.item_type, "status": item.status, "priority": item.priority,
        "assigned_role": item.assigned_role, "estimate_points": item.estimate_points, "epic_id": item.project_epic_id,
        "assignee_id": item.assignee_id, "sprint_id": item.sprint_id, "due_date": item.due_date,
        "order_index": item.order_index,
    }
    for name, value in data.items():
        if name == "description" or name == "acceptance_criteria":
            value = (value or "").strip()
        setattr(item, name, value)
    if "status" in data:
        item.completed_at = datetime.now(timezone.utc) if item.status == "DONE" else None
    after = {
        "title": item.title, "description": item.description, "acceptance_criteria": item.acceptance_criteria,
        "item_type": item.item_type, "status": item.status, "priority": item.priority,
        "assigned_role": item.assigned_role, "estimate_points": item.estimate_points, "epic_id": item.project_epic_id,
        "assignee_id": item.assignee_id, "sprint_id": item.sprint_id, "due_date": item.due_date,
        "order_index": item.order_index,
    }
    for field, old in before.items():
        if old != after[field]:
            await _add_history(db, item, actor, field, old, after[field])
    if before["sprint_id"] != item.sprint_id:
        if before["sprint_id"] is not None:
            prior_link = await db.scalar(select(AgileSprintItem).where(
                AgileSprintItem.sprint_id == before["sprint_id"], AgileSprintItem.work_item_id == item.id
            ))
            if prior_link is not None and prior_link.outcome is None:
                prior_link.outcome = "MOVED_TO_SPRINT" if item.sprint_id is not None else "MOVED_TO_BACKLOG"
                prior_link.closed_at = datetime.now(timezone.utc)
        if item.sprint_id is not None:
            next_link = await db.scalar(select(AgileSprintItem).where(
                AgileSprintItem.sprint_id == item.sprint_id, AgileSprintItem.work_item_id == item.id
            ))
            if next_link is not None and next_link.outcome is not None:
                next_link.outcome = None
                next_link.closed_at = None
            await _record_sprint_item(db, item.sprint_id, item.id)
    await db.commit()
    return await _item_read(db, item)


@router.delete("/projects/{project_id}/items/{item_id}", status_code=204)
async def archive_item(project_id: int, item_id: int, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    item = await db.get(AgileWorkItem, item_id)
    if item is None or item.project_id != project_id or item.is_archived:
        raise AppError(404, "Tarjeta no encontrada.")
    item.is_archived = True
    await _add_history(db, item, actor, "archived", False, True)
    await db.commit()
    return None


@router.get("/projects/{project_id}/sprints", response_model=list[SprintRead])
async def list_sprints(project_id: int, db: DBSession, actor: StaffUser):
    await _project_access(db, project_id, actor)
    sprints = (await db.scalars(select(AgileSprint).where(AgileSprint.project_id == project_id).order_by(AgileSprint.start_date, AgileSprint.id))).all()
    return [await _sprint_read(db, sprint) for sprint in sprints]


@router.post("/projects/{project_id}/sprints", response_model=SprintRead, status_code=201)
async def create_sprint(project_id: int, payload: SprintCreate, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    if await db.scalar(select(AgileSprint.id).where(
        AgileSprint.project_id == project_id, AgileSprint.name == payload.name.strip()
    )) is not None:
        raise AppError(409, "Ya existe un sprint con ese nombre en el proyecto.")
    sprint = AgileSprint(project_id=project_id, name=payload.name.strip(), goal=payload.goal.strip(),
                         start_date=payload.start_date, end_date=payload.end_date,
                         status="PLANNED", created_by_id=actor.id)
    db.add(sprint)
    await db.commit()
    await db.refresh(sprint)
    return await _sprint_read(db, sprint)


@router.post("/projects/{project_id}/sprints/{sprint_id}/start", response_model=SprintRead)
async def start_sprint(project_id: int, sprint_id: int, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    sprint = await db.get(AgileSprint, sprint_id)
    if sprint is None or sprint.project_id != project_id:
        raise AppError(404, "Sprint no encontrado.")
    if sprint.status != "PLANNED":
        raise AppError(409, "Solo se puede iniciar un sprint planificado.")
    active = await db.scalar(select(AgileSprint).where(
        AgileSprint.project_id == project_id, AgileSprint.status == "ACTIVE"
    ))
    if active:
        raise AppError(409, f"Ya existe un sprint activo: {active.name}.")
    sprint.status = "ACTIVE"
    await db.commit()
    return await _sprint_read(db, sprint)


@router.post("/projects/{project_id}/sprints/{sprint_id}/complete", response_model=dict)
async def complete_sprint(project_id: int, sprint_id: int, payload: SprintComplete, db: DBSession, actor: StaffUser):
    project, _ = await _require_role(db, project_id, actor, PLAN_ROLES)
    sprint = await db.get(AgileSprint, sprint_id)
    if sprint is None or sprint.project_id != project_id:
        raise AppError(404, "Sprint no encontrado.")
    if sprint.status != "ACTIVE":
        raise AppError(409, "Solo se puede completar un sprint activo.")
    next_sprint: AgileSprint | None = None
    now = datetime.now(timezone.utc)
    items = (await db.scalars(select(AgileWorkItem).where(
        AgileWorkItem.project_id == project_id, AgileWorkItem.sprint_id == sprint_id,
        AgileWorkItem.is_archived.is_(False)
    ))).all()
    has_pending_items = any(item.status != "DONE" for item in items)
    if payload.carry_over and payload.next_sprint_id:
        next_sprint = await db.get(AgileSprint, payload.next_sprint_id)
        if next_sprint is None or next_sprint.project_id != project_id or next_sprint.status != "PLANNED" or next_sprint.id == sprint.id:
            raise AppError(422, "El sprint de destino debe ser un sprint planificado del mismo proyecto.")
    elif payload.carry_over and has_pending_items:
        next_sprint = await db.scalar(select(AgileSprint).where(
            AgileSprint.project_id == project_id, AgileSprint.status == "PLANNED",
            AgileSprint.start_date >= sprint.start_date
        ).order_by(AgileSprint.start_date, AgileSprint.id))
        if next_sprint is None:
            sprint_count = await db.scalar(select(func.count(AgileSprint.id)).where(AgileSprint.project_id == project_id)) or 0
            start = sprint.end_date + timedelta(days=1)
            next_sprint = AgileSprint(
                project_id=project_id, name=f"Sprint {sprint_count + 1}", goal="",
                start_date=start, end_date=start + timedelta(weeks=project.sprint_length_weeks) - timedelta(days=1),
                status="PLANNED", created_by_id=actor.id,
            )
            db.add(next_sprint)
            await db.flush()
    carried = returned = done = 0
    for item in items:
        link = await db.scalar(select(AgileSprintItem).where(
            AgileSprintItem.sprint_id == sprint_id, AgileSprintItem.work_item_id == item.id
        ))
        if item.status == "DONE":
            done += 1
            if link:
                link.outcome, link.closed_at = "COMPLETED", now
            item.sprint_id = None
            await _add_history(db, item, actor, "sprint_id", sprint_id, None)
        elif payload.carry_over:
            carried += 1
            if link:
                link.outcome, link.closed_at = "CARRIED_OVER", now
            item.sprint_id = next_sprint.id
            await _record_sprint_item(db, next_sprint.id, item.id)
            await _add_history(db, item, actor, "sprint_id", sprint_id, next_sprint.id)
        else:
            returned += 1
            if link:
                link.outcome, link.closed_at = "RETURNED_TO_BACKLOG", now
            item.sprint_id = None
            await _add_history(db, item, actor, "sprint_id", sprint_id, None)
    sprint.status = "COMPLETED"
    sprint.completed_at = now
    await db.commit()
    return {
        "completed_sprint": await _sprint_read(db, sprint),
        "next_sprint": await _sprint_read(db, next_sprint) if next_sprint else None,
        "completed_items": done, "carried_over": carried, "returned_to_backlog": returned,
    }


# ---------- importación XLSX / CSV con vista previa, límites y validación antes del commit
_HEADER_ALIASES = {
    "title": {"title", "titulo", "título", "nombre", "summary", "resumen"},
    "type": {"type", "tipo", "item_type", "issue type", "issue_type"},
    "description": {"description", "descripcion", "descripción", "detalle"},
    "epic": {"epic", "epica", "épica", "epic key", "épica clave", "epic_key", "epic title"},
    "priority": {"priority", "prioridad"},
    "assigned_role": {"assigned_role", "assigned role", "rol responsable", "rol asignado", "responsable rol"},
    "status": {"status", "estado"},
    "points": {"points", "puntos", "story points", "story_points", "estimacion", "estimación"},
    "assignee_email": {"assignee_email", "responsable_email", "email responsable", "assigned to", "asignado a", "email"},
    "sprint": {"sprint", "sprint_name", "iteracion", "iteración"},
    "acceptance_criteria": {"acceptance_criteria", "criterios de aceptacion", "criterios de aceptación", "criterios"},
    "due_date": {"due_date", "fecha límite", "fecha limite", "fecha vencimiento"},
}


def _norm_header(raw: Any) -> str:
    value = unicodedata.normalize("NFKC", str(raw or "")).strip().casefold()
    return re.sub(r"\s+", " ", value)


def _normalise_type(value: Any) -> str:
    text_value = _norm_header(value).upper()
    aliases = {"EPICA": "EPIC", "ÉPICA": "EPIC", "HISTORIA": "STORY", "HISTORIA DE USUARIO": "STORY",
               "HU": "STORY", "TAREA": "TASK", "BUG": "BUG", "ERROR": "BUG", "SUBTAREA": "SUBTASK"}
    return aliases.get(text_value, text_value or "STORY")


def _normalise_priority(value: Any) -> str:
    text_value = _norm_header(value).upper()
    aliases = {"BAJA": "LOW", "LOW": "LOW", "MEDIA": "MEDIUM", "MEDIUM": "MEDIUM",
               "ALTA": "HIGH", "HIGH": "HIGH", "URGENTE": "URGENT", "CRÍTICA": "URGENT", "CRITICA": "URGENT"}
    return aliases.get(text_value, text_value or "MEDIUM")


def _normalise_status(value: Any) -> str:
    text_value = _norm_header(value).upper().replace(" ", "_")
    aliases = {"PENDIENTE": "TODO", "POR_HACER": "TODO", "TODO": "TODO", "EN_PROGRESO": "IN_PROGRESS",
               "EN_CURSO": "IN_PROGRESS", "IN_PROGRESS": "IN_PROGRESS", "REVISION": "IN_REVIEW",
               "EN_REVISION": "IN_REVIEW", "IN_REVIEW": "IN_REVIEW", "HECHO": "DONE", "COMPLETADO": "DONE",
               "DONE": "DONE", "BLOQUEADO": "BLOCKED", "BLOCKED": "BLOCKED"}
    return aliases.get(text_value, text_value or "TODO")


def _normalise_assigned_role(value: Any) -> str:
    text_value = _norm_header(value).upper().replace(" ", "_").replace("-", "_")
    aliases = {
        "SM": "SCRUM_MASTER", "SCRUMMASTER": "SCRUM_MASTER", "SCRUM_MASTER": "SCRUM_MASTER",
        "PO": "PRODUCT_OWNER", "PRODUCTOWNER": "PRODUCT_OWNER", "PRODUCT_OWNER": "PRODUCT_OWNER",
        "DEV": "DEVELOPER", "DEVELOPER": "DEVELOPER", "DESARROLLADOR": "DEVELOPER",
        "SCRUM_MASTER": "SCRUM_MASTER", "PRODUCT_OWNER": "PRODUCT_OWNER",
    }
    return aliases.get(text_value, text_value)


async def _read_import(file: UploadFile) -> tuple[str, list[dict[str, Any]]]:
    filename = Path(file.filename or "import.csv").name
    data = await file.read(MAX_IMPORT_BYTES + 1)
    if len(data) > MAX_IMPORT_BYTES:
        raise AppError(413, "El archivo supera el límite de 5 MB.")
    ext = Path(filename).suffix.lower()
    raw_rows: list[list[Any]] = []
    try:
        if ext == ".csv":
            text_data = data.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text_data[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            raw_rows = list(csv.reader(io.StringIO(text_data), dialect))
        elif ext in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            worksheet = workbook.active
            raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            workbook.close()
        else:
            raise AppError(415, "Formato no compatible. Subí un archivo .xlsx o .csv.")
    except AppError:
        raise
    except Exception:
        raise AppError(422, "No se pudo leer el archivo. Verificá que sea un Excel o CSV válido.") from None
    if not raw_rows:
        raise AppError(422, "El archivo no contiene filas.")
    raw_headers = [_norm_header(h) for h in raw_rows[0]]
    mapping: dict[str, int] = {}
    for i, header in enumerate(raw_headers):
        for field, aliases in _HEADER_ALIASES.items():
            if header in aliases and field not in mapping:
                mapping[field] = i
    if "title" not in mapping:
        raise AppError(422, "Falta una columna de título. Usá, por ejemplo, 'Título' o 'Title'.")
    rows: list[dict[str, Any]] = []
    for number, raw in enumerate(raw_rows[1:], start=2):
        if not any(cell is not None and str(cell).strip() for cell in raw):
            continue
        def value(field: str) -> str:
            idx = mapping.get(field)
            return str(raw[idx]).strip() if idx is not None and idx < len(raw) and raw[idx] is not None else ""
        points_raw = value("points")
        points: int | None = None
        errors: list[str] = []
        if points_raw:
            try:
                points_num = float(points_raw.replace(",", "."))
                if points_num < 0 or not points_num.is_integer() or points_num > 1000:
                    raise ValueError
                points = int(points_num)
            except ValueError:
                errors.append("Los puntos deben ser un entero entre 0 y 1000.")
        item_type = _normalise_type(value("type"))
        status = _normalise_status(value("status"))
        priority = _normalise_priority(value("priority"))
        assigned_role = _normalise_assigned_role(value("assigned_role")) if value("assigned_role") else ""
        if item_type not in ITEM_TYPES:
            errors.append(f"Tipo no reconocido: {item_type}.")
        if status not in STATUSES:
            errors.append(f"Estado no reconocido: {status}.")
        if priority not in PRIORITIES:
            errors.append(f"Prioridad no reconocida: {priority}.")
        if assigned_role and assigned_role not in {"SCRUM_MASTER", "PRODUCT_OWNER", "DEVELOPER"}:
            errors.append(f"Rol responsable no reconocido: {assigned_role}.")
        title = value("title")
        if not title:
            errors.append("El título es obligatorio.")
        due_date_raw = raw[mapping["due_date"]] if "due_date" in mapping and mapping["due_date"] < len(raw) else None
        parsed_due_date: date | None = None
        if due_date_raw not in (None, ""):
            try:
                if isinstance(due_date_raw, datetime):
                    parsed_due_date = due_date_raw.date()
                elif isinstance(due_date_raw, date):
                    parsed_due_date = due_date_raw
                else:
                    raw_date_text = str(due_date_raw).strip()
                    try:
                        parsed_due_date = date.fromisoformat(raw_date_text[:10])
                    except ValueError:
                        parsed_due_date = None
                        for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
                            try:
                                parsed_due_date = datetime.strptime(raw_date_text, fmt).date()
                                break
                            except ValueError:
                                continue
                    if parsed_due_date is None:
                        raise ValueError
            except (ValueError, TypeError):
                errors.append("La fecha límite debe tener formato AAAA-MM-DD o DD/MM/AAAA.")
        rows.append({
            "row_number": number, "title": title, "item_type": item_type, "description": value("description"),
            "epic": value("epic"), "priority": priority, "status": status, "points": points,
            "assigned_role": assigned_role, "assignee_email": value("assignee_email").casefold(), "sprint": value("sprint"),
            "acceptance_criteria": value("acceptance_criteria"), "due_date": parsed_due_date, "errors": errors,
        })
        if len(rows) > MAX_IMPORT_ROWS:
            raise AppError(413, f"El archivo supera el máximo de {MAX_IMPORT_ROWS} filas de datos.")
    return filename, rows


async def _preview_import(db: AsyncSession, project_id: int, filename: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    existing_epics = (await db.scalars(select(AgileWorkItem).where(
        AgileWorkItem.project_id == project_id, AgileWorkItem.item_type == "EPIC", AgileWorkItem.is_archived.is_(False)
    ))).all()
    epic_refs = {epic.title.casefold() for epic in existing_epics} | {epic.key.casefold() for epic in existing_epics}
    import_epics = {r["title"].casefold() for r in rows if r["item_type"] == "EPIC" and r["title"]}
    members = (await db.execute(
        select(User.email).join(AgileMembership, AgileMembership.user_id == User.id)
        .where(AgileMembership.project_id == project_id)
    )).scalars().all()
    member_emails = {email.casefold() for email in members}
    sprints = (await db.scalars(select(AgileSprint).where(
        AgileSprint.project_id == project_id, AgileSprint.status != "COMPLETED"
    ))).all()
    sprint_names = {s.name.casefold() for s in sprints}
    for row in rows:
        if row["epic"] and row["item_type"] != "EPIC" and row["epic"].casefold() not in epic_refs | import_epics:
            row["errors"].append(f"No se encontró la épica '{row['epic']}'. Importala o corregí su clave/título.")
        if row["assignee_email"] and row["assignee_email"] not in member_emails:
            row["errors"].append(f"El responsable {row['assignee_email']} no es integrante del proyecto.")
        if row["sprint"] and row["sprint"].casefold() not in sprint_names:
            row["errors"].append(f"El sprint '{row['sprint']}' no existe o ya está completado.")
        if row["item_type"] == "EPIC" and row["epic"]:
            row["errors"].append("Una épica no puede estar subordinada a otra épica.")
    valid_count = sum(1 for r in rows if not r["errors"])
    shown = rows[:100]
    return {
        "filename": filename, "total_rows": len(rows), "valid_rows": valid_count,
        "error_rows": sum(1 for r in rows if r["errors"]),
        "rows": [{k: r[k] for k in ("row_number", "title", "item_type", "description", "epic", "priority", "status", "points", "assigned_role", "assignee_email", "sprint", "acceptance_criteria", "due_date", "errors")} for r in shown],
        "truncated": len(rows) > len(shown),
    }


def _sheet_export_url(user_url: str) -> str:
    parts = urlsplit(user_url.strip())
    if parts.scheme != "https" or parts.hostname not in {"docs.google.com", "drive.google.com"}:
        raise AppError(422, "Pegá un enlace HTTPS de Google Sheets (docs.google.com/spreadsheets/d/…).")
    match = re.fullmatch(r"/spreadsheets/d/([A-Za-z0-9_-]{10,})/(?:.*)?", parts.path)
    if not match:
        raise AppError(422, "El enlace no parece ser de una hoja de cálculo de Google Sheets.")
    spreadsheet_id = match.group(1)
    gid = parse_qs(parts.query).get("gid", ["0"])[0]
    if not re.fullmatch(r"[0-9]+", gid):
        gid = "0"
    return f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=csv&gid={gid}"


async def _google_sheet_upload(user_url: str) -> StarletteUploadFile:
    """Descarga solo exportaciones CSV de Google Sheets; nunca acepta una URL arbitraria."""
    import httpx

    current = _sheet_export_url(user_url)
    timeout = httpx.Timeout(12.0, connect=5.0)
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=False) as client:
        for _ in range(5):
            try:
                response = await client.get(current)
            except httpx.HTTPError:
                raise AppError(502, "No se pudo conectar con Google Sheets. Intentá exportar la hoja como CSV.") from None
            if response.status_code in {301, 302, 303, 307, 308}:
                location = response.headers.get("location")
                if not location:
                    raise AppError(502, "Google Sheets devolvió una redirección inválida.")
                next_url = urljoin(current, location)
                parts = urlsplit(next_url)
                host = parts.hostname or ""
                if parts.scheme != "https" or not (host in {"docs.google.com", "drive.google.com"} or host.endswith(".googleusercontent.com")):
                    raise AppError(422, "Google Sheets redirigió a un dominio no permitido.")
                current = next_url
                continue
            if response.status_code in {401, 403}:
                raise AppError(403, "La hoja no es accesible. Compartila con acceso de lectura por enlace o importá un CSV/XLSX descargado.")
            if response.status_code >= 400:
                raise AppError(422, "Google Sheets no pudo exportar la hoja. Revisá el enlace y sus permisos.")
            content = response.content
            if len(content) > MAX_IMPORT_BYTES:
                raise AppError(413, "La hoja exportada supera el límite de 5 MB.")
            content_type = response.headers.get("content-type", "").lower()
            if "html" in content_type or content.lstrip().lower().startswith((b"<!doctype html", b"<html")):
                raise AppError(403, "Google Sheets devolvió una página de acceso. Habilitá lectura por enlace o subí el archivo exportado.")
            if not content.strip():
                raise AppError(422, "La hoja está vacía o no contiene datos exportables.")
            return StarletteUploadFile(filename="google-sheets.csv", file=io.BytesIO(content))
    raise AppError(502, "Google Sheets realizó demasiadas redirecciones.")


@router.post("/projects/{project_id}/imports/google/preview", response_model=ImportPreview)
async def preview_google_import(project_id: int, payload: GoogleSheetImportRequest, db: DBSession, actor: StaffUser):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    file = await _google_sheet_upload(payload.url)
    filename, rows = await _read_import(file)
    return await _preview_import(db, project_id, filename, rows)


@router.post("/projects/{project_id}/imports/google/commit", response_model=ImportResult)
async def commit_google_import(
    project_id: int, payload: GoogleSheetImportRequest, db: DBSession, actor: StaffUser,
    duplicate_strategy: str = Query(default="skip", pattern="^(skip|update)$"),
):
    file = await _google_sheet_upload(payload.url)
    return await commit_import(project_id, db, actor, file, duplicate_strategy=duplicate_strategy)


@router.post("/projects/{project_id}/imports/preview", response_model=ImportPreview)
async def preview_import(project_id: int, db: DBSession, actor: StaffUser, file: Annotated[UploadFile, File()]):
    await _require_role(db, project_id, actor, PLAN_ROLES)
    filename, rows = await _read_import(file)
    return await _preview_import(db, project_id, filename, rows)


@router.post("/projects/{project_id}/imports/commit", response_model=ImportResult)
async def commit_import(
    project_id: int, db: DBSession, actor: StaffUser, file: Annotated[UploadFile, File()],
    duplicate_strategy: str = Query(default="skip", pattern="^(skip|update)$"),
):
    project, _ = await _require_role(db, project_id, actor, PLAN_ROLES)
    filename, rows = await _read_import(file)
    preview = await _preview_import(db, project_id, filename, rows)
    if preview["error_rows"]:
        raise AppError(422, {"message": "La importación contiene errores; no se guardó ningún registro.", "preview": preview})
    existing = (await db.scalars(select(AgileWorkItem).where(
        AgileWorkItem.project_id == project_id, AgileWorkItem.is_archived.is_(False)
    ))).all()
    # Se detectan duplicados por tipo + título exacto (sin distinguir mayúsculas). No se fusionan HU distintas por aproximación.
    index = {(item.item_type, item.title.casefold()): item for item in existing}
    member_by_email_rows = (await db.execute(
        select(User.email, User.id).join(AgileMembership, AgileMembership.user_id == User.id)
        .where(AgileMembership.project_id == project_id)
    )).all()
    member_by_email = {email.casefold(): user_id for email, user_id in member_by_email_rows}
    sprints = (await db.scalars(select(AgileSprint).where(
        AgileSprint.project_id == project_id, AgileSprint.status != "COMPLETED"
    ))).all()
    sprint_by_name = {s.name.casefold(): s for s in sprints}
    epic_by_ref: dict[str, AgileWorkItem] = {}
    for item in existing:
        if item.item_type == "EPIC":
            epic_by_ref[item.title.casefold()] = item
            epic_by_ref[item.key.casefold()] = item

    # Crear primero las épicas para permitir que las HU de la misma hoja las referencien.
    created = updated = skipped = 0
    now = datetime.now(timezone.utc)
    ordered_rows = [r for r in rows if r["item_type"] == "EPIC"] + [r for r in rows if r["item_type"] != "EPIC"]
    for row in ordered_rows:
        title = row["title"].strip()
        key = (row["item_type"], title.casefold())
        if key in index:
            if duplicate_strategy == "skip":
                skipped += 1
                continue
            item = index[key]
            old_title = item.title
            item.description = row["description"]
            item.acceptance_criteria = row["acceptance_criteria"]
            item.priority = row["priority"]
            item.assigned_role = row["assigned_role"] or None
            item.status = row["status"]
            item.estimate_points = row["points"]
            item.due_date = row["due_date"]
            item.completed_at = now if item.status == "DONE" else None
            await _add_history(db, item, actor, "import_updated", old_title, item.title)
            updated += 1
        else:
            item = AgileWorkItem(
                project_id=project_id, key=f"TMP-{uuid4().hex[:12].upper()}", title=title, description=row["description"],
                acceptance_criteria=row["acceptance_criteria"], item_type=row["item_type"],
                status=row["status"], priority=row["priority"], assigned_role=row["assigned_role"] or None,
                estimate_points=row["points"], due_date=row["due_date"], created_by_id=actor.id, completed_at=now if row["status"] == "DONE" else None,
            )
            db.add(item)
            await db.flush()
            item.key = f"{project.key}-{item.id:04d}"
            index[key] = item
            created += 1
        if item.item_type == "EPIC":
            epic_by_ref[item.title.casefold()] = item
            epic_by_ref[item.key.casefold()] = item
        assignee_email = row["assignee_email"]
        if assignee_email:
            item.assignee_id = member_by_email[assignee_email]
        if row["sprint"]:
            sprint = sprint_by_name[row["sprint"].casefold()]
            item.sprint_id = sprint.id
            await _record_sprint_item(db, sprint.id, item.id)
        elif item.item_type != "EPIC" and item.sprint_id and item.id not in {e.id for e in existing}:
            item.sprint_id = None
        if row["item_type"] != "EPIC" and row["epic"]:
            epic = epic_by_ref.get(row["epic"].casefold())
            if epic is None:
                # El preflight evita esta rama: conservar rollback si la hoja cambia durante el proceso.
                await db.rollback()
                raise AppError(422, f"No se pudo resolver la épica '{row['epic']}'. No se guardó la importación.")
            item.project_epic_id = epic.id
        db.add(AgileWorkItemHistory(project_id=project_id, work_item_id=item.id, actor_id=actor.id,
                                    field_name="imported", old_value=None, new_value=filename))
    await db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "errors": []}


@router.get("/projects/{project_id}/items/{item_id}/history", response_model=list[WorkItemHistoryRead])
async def list_work_item_history(project_id: int, item_id: int, db: DBSession, actor: StaffUser):
    await _project_access(db, project_id, actor)
    item = await db.get(AgileWorkItem, item_id)
    if item is None or item.project_id != project_id:
        raise AppError(404, "Tarjeta no encontrada.")
    rows = (await db.execute(
        select(AgileWorkItemHistory, User)
        .join(User, User.id == AgileWorkItemHistory.actor_id)
        .where(
            AgileWorkItemHistory.project_id == project_id,
            AgileWorkItemHistory.work_item_id == item_id,
        )
        .order_by(AgileWorkItemHistory.created_at.desc(), AgileWorkItemHistory.id.desc())
        .limit(500)
    )).all()
    return [{
        "id": entry.id,
        "work_item_id": entry.work_item_id,
        "actor_id": entry.actor_id,
        "actor_name": user.full_name,
        "field_name": entry.field_name,
        "old_value": entry.old_value,
        "new_value": entry.new_value,
        "created_at": entry.created_at,
    } for entry, user in rows]


@router.get("/projects/{project_id}/items/{item_id}/comments", response_model=list[WorkItemCommentRead])
async def list_work_item_comments(project_id: int, item_id: int, db: DBSession, actor: StaffUser):
    await _project_access(db, project_id, actor)
    item = await db.get(AgileWorkItem, item_id)
    if item is None or item.project_id != project_id:
        raise AppError(404, "Tarjeta no encontrada.")
    rows = (await db.execute(
        select(AgileWorkItemComment, User)
        .join(User, User.id == AgileWorkItemComment.author_id)
        .where(
            AgileWorkItemComment.project_id == project_id,
            AgileWorkItemComment.work_item_id == item_id,
        )
        .order_by(AgileWorkItemComment.created_at.asc(), AgileWorkItemComment.id.asc())
        .limit(1000)
    )).all()
    return [{
        "id": comment.id,
        "work_item_id": comment.work_item_id,
        "author_id": comment.author_id,
        "author_name": user.full_name,
        "author_email": user.email,
        "content": comment.content,
        "created_at": comment.created_at,
        "updated_at": comment.updated_at,
    } for comment, user in rows]


@router.post("/projects/{project_id}/items/{item_id}/comments", response_model=WorkItemCommentRead, status_code=201)
async def create_work_item_comment(
    project_id: int, item_id: int, payload: WorkItemCommentCreate, db: DBSession, actor: StaffUser
):
    _, role = await _project_access(db, project_id, actor)
    if role == "OBSERVER":
        raise AppError(403, "Los observadores pueden consultar las tarjetas, pero no comentar.")
    item = await db.get(AgileWorkItem, item_id)
    if item is None or item.project_id != project_id or item.is_archived:
        raise AppError(404, "Tarjeta no encontrada.")
    content = payload.content.strip()
    if not content:
        raise AppError(422, "El comentario no puede quedar vacío.")
    now = datetime.now(timezone.utc)
    comment = AgileWorkItemComment(
        project_id=project_id,
        work_item_id=item_id,
        author_id=actor.id,
        content=content,
        updated_at=now,
    )
    db.add(comment)
    await db.flush()
    db.add(AgileWorkItemHistory(
        project_id=project_id,
        work_item_id=item_id,
        actor_id=actor.id,
        field_name="commented",
        old_value=None,
        new_value=content[:200],
    ))
    await db.commit()
    await db.refresh(comment)
    return {
        "id": comment.id,
        "work_item_id": comment.work_item_id,
        "author_id": comment.author_id,
        "author_name": actor.full_name,
        "author_email": actor.email,
        "content": comment.content,
        "created_at": comment.created_at,
        "updated_at": comment.updated_at,
    }
